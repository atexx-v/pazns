from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Arial"
MONO = "Menlo"
ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}
TOTAL_WIDTH_DXA = 9638

def _set_run_font(run, name, size, bold=False, italic=False, color=None):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        from docx.shared import RGBColor
        run.font.color.rgb = RGBColor.from_string(color)

def _spacing(par, before=0, after=0, line=None):
    fmt = par.paragraph_format
    fmt.space_before = Pt(before / 20)
    fmt.space_after = Pt(after / 20)
    if line:
        ppr = par._p.get_or_add_pPr()
        sp = ppr.find(qn("w:spacing"))
        sp.set(qn("w:line"), str(line))
        sp.set(qn("w:lineRule"), "auto")

def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(14)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    return doc

def para(doc, text="", size=14, bold=False, italic=False, align="left",
         before=0, after=0, line=240, keep_next=False, font=FONT, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.alignment = ALIGN[align]
    _spacing(p, before, after, line)
    p.paragraph_format.keep_with_next = keep_next
    if text:
        _set_run_font(p.add_run(text), font, size, bold, italic)
    return p

def rich(doc, parts, size=14, align="left", before=0, after=120, line=240, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.alignment = ALIGN[align]
    _spacing(p, before, after, line)
    for text, fmt in parts:
        _set_run_font(p.add_run(text), FONT, size, fmt.get("bold", False), fmt.get("italic", False))
    return p

def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def title_page(doc, number, logo_path, group="ІР-33", student="ВИШНЕВСЬКИЙ А. Г.",
               teacher="ШИХМАТ А.О", year="2026",
               subject="З ДИСЦИПЛІНИ «ПРОГРАМНО-АПАРАТНЕ ЗАБЕЗПЕЧЕННЯ СИСТЕМ РУХОМИХ ОБ'ЄКТІВ»"):
    para(doc, "МІНІСТЕРСТВО ОСВІТИ І НАУКИ", 16, bold=True, align="center", line=360)
    para(doc, "НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ «ЛЬВІВСЬКА ПОЛІТЕХНІКА»", 16, bold=True, align="center",
         after=360, line=360)
    logo = doc.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo.add_run().add_picture(logo_path, width=Cm(7.0))
    para(doc, after=480)
    para(doc, f"ЛАБОРАТОРНА РОБОТА № {number}", 14, align="center", after=200, line=360)
    para(doc, subject, 12, italic=True, align="center")
    for _ in range(3):
        para(doc, after=280)
    para(doc, "ВИКОНАВ:", 14, bold=True, align="right", line=360)
    para(doc, f"{student} ({group})", 14, italic=True, align="right", line=360)
    para(doc, "ПРИЙНЯВ:", 14, bold=True, align="right", line=360)
    para(doc, teacher, 14, align="right", line=360)
    for _ in range(2):
        para(doc, after=280)
    para(doc, f"ЛЬВІВ {year}", 16, bold=True, italic=True, align="center", line=360)
    page_break(doc)

def heading(doc, text):
    return para(doc, text, 14, bold=True, italic=True, before=240, after=100, keep_next=True)

def subheading(doc, text):
    return para(doc, text, 14, bold=True, before=160, after=120, keep_next=True)

def text(doc, content, after=120):
    return para(doc, content, 14, after=after)

def bullet(doc, content):
    return para(doc, content, 14, after=60, style="List Bullet")

def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)

def table(doc, rows, widths, caption=None, header=True, size=11):
    if caption:
        para(doc, caption, 12, align="right", after=60, keep_next=True)
    tbl = doc.add_table(rows=len(rows), cols=len(widths))
    tblpr = tbl._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblpr.append(borders)
    grid = tbl._tbl.tblGrid
    for col, w in zip(grid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(w))
    for r, row in enumerate(rows):
        tr = tbl.rows[r]._tr
        tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for c, value in enumerate(row):
            cell = tbl.rows[r].cells[c]
            tcw = cell._tc.get_or_add_tcPr()
            w = OxmlElement("w:tcW")
            w.set(qn("w:type"), "dxa")
            w.set(qn("w:w"), str(widths[c]))
            tcw.append(w)
            if header and r == 0:
                _shade(cell, "D9D9D9")
            p = cell.paragraphs[0]
            _spacing(p, 20, 20, 240)
            _set_run_font(p.add_run(str(value)), FONT, size, bold=(header and r == 0))
    para(doc, after=120, line=240)
    return tbl

def code_block(doc, content, size=7.5):
    lines = content.rstrip("\n").split("\n")
    p = doc.add_paragraph()
    _spacing(p, 120, 280, 240)
    ppr = p._p.get_or_add_pPr()
    ppr.append(OxmlElement("w:keepLines"))
    pbdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "6")
        el.set(qn("w:color"), "000000")
        pbdr.append(el)
    ppr.append(pbdr)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "000000")
    ppr.append(shd)
    for i, line in enumerate(lines):
        stripped = line.lstrip(" ")
        indent = len(line) - len(stripped)
        run = p.add_run("  " + " " * indent + stripped)
        _set_run_font(run, MONO, size, color="F0F0F0")
        if i < len(lines) - 1:
            run.add_break()
    return p

def figure(doc, path, width_cm, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p, 0, 60)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width_cm))
    para(doc, caption, 12, align="center", after=240)
